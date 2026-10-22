from pathlib import Path
import re
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    print("\nERROR: Falta la librería 'openpyxl'.")
    print("Instálala una sola vez con este comando:")
    print("    py -m pip install openpyxl")
    input("\nPresiona Enter para cerrar...")
    raise SystemExit(1)


COLUMNA_LLAVE = "LLAVE"
CARACTERES_INVALIDOS = re.compile(r'[<>:"/\\|?*\x00-\x1F]')
NOMBRES_RESERVADOS = {
    "CON", "PRN", "AUX", "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}


def limpiar_nombre_carpeta(valor) -> str:
    """Convierte una llave del Excel en un nombre válido de carpeta para Windows."""
    nombre = str(valor).strip()

    # En el archivo las llaves usan | como separador, pero Windows no permite |.
    nombre = nombre.replace("|", " - ")

    # Reemplaza cualquier otro carácter no permitido en nombres de carpetas.
    nombre = CARACTERES_INVALIDOS.sub("_", nombre)

    # Limpieza visual.
    nombre = re.sub(r"\s+", " ", nombre).strip(" .")

    # Windows tampoco permite estos nombres reservados.
    if nombre.upper() in NOMBRES_RESERVADOS:
        nombre = f"_{nombre}"

    return nombre


def buscar_excel(carpeta: Path) -> Path:
    """Busca el Excel a leer en la carpeta donde se ejecutó el programa."""
    archivos = sorted(
        p for p in carpeta.glob("*.xlsx")
        if not p.name.startswith("~$")
    )

    if not archivos:
        print("\nNo encontré ningún archivo .xlsx en esta carpeta:")
        print(f"  {carpeta}")
        print("\nColoca este programa junto al Excel y vuelve a ejecutarlo.")
        raise SystemExit(1)

    if len(archivos) == 1:
        return archivos[0]

    print("\nEncontré varios archivos de Excel. Elige cuál contiene las llaves:\n")
    for i, archivo in enumerate(archivos, start=1):
        print(f"  {i}. {archivo.name}")

    while True:
        try:
            opcion = int(input("\nNúmero de archivo: "))
            if 1 <= opcion <= len(archivos):
                return archivos[opcion - 1]
        except ValueError:
            pass
        print("Opción no válida. Intenta nuevamente.")


def localizar_columna_llave(ws):
    """Busca la cabecera LLAVE en las primeras 20 filas de la hoja."""
    for fila in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
        for celda in fila:
            if str(celda.value).strip().upper() == COLUMNA_LLAVE:
                return celda.row, celda.column
    return None


def obtener_llaves(archivo_excel: Path):
    """Obtiene las llaves de la primera hoja que contenga la columna LLAVE."""
    wb = load_workbook(archivo_excel, read_only=True, data_only=True)

    try:
        for ws in wb.worksheets:
            posicion = localizar_columna_llave(ws)
            if posicion is None:
                continue

            fila_encabezado, columna = posicion
            llaves = []

            for fila in ws.iter_rows(
                min_row=fila_encabezado + 1,
                min_col=columna,
                max_col=columna,
                values_only=True,
            ):
                valor = fila[0]
                if valor is None:
                    continue
                texto = str(valor).strip()
                if texto:
                    llaves.append(texto)

            return ws.title, llaves
    finally:
        wb.close()

    raise ValueError(
        f"No encontré una columna llamada '{COLUMNA_LLAVE}' en ninguna hoja del Excel."
    )


def crear_carpetas(llaves, carpeta_destino: Path):
    creadas = 0
    existentes = 0
    omitidas = 0
    vistos = set()

    for llave in llaves:
        nombre = limpiar_nombre_carpeta(llave)

        if not nombre:
            omitidas += 1
            continue

        # Evita procesar dos veces la misma llave repetida en el Excel.
        clave_comparacion = nombre.casefold()
        if clave_comparacion in vistos:
            omitidas += 1
            continue
        vistos.add(clave_comparacion)

        ruta = carpeta_destino / nombre

        if ruta.exists():
            existentes += 1
        else:
            ruta.mkdir(parents=False, exist_ok=True)
            creadas += 1
            print(f"[CREADA] {nombre}")

    return creadas, existentes, omitidas


def main():
    carpeta_destino = Path.cwd()

    print("=" * 64)
    print("  CREADOR DE CARPETAS POR LLAVE DE VEHÍCULO")
    print("=" * 64)
    print(f"\nCarpeta de destino:\n  {carpeta_destino}")

    # También permite pasar el Excel manualmente:
    # py crear_carpetas_llaves.py "C:\\ruta\\archivo.xlsx"
    if len(sys.argv) > 1:
        archivo_excel = Path(sys.argv[1]).expanduser().resolve()
        if not archivo_excel.is_file():
            print(f"\nERROR: No existe el archivo:\n  {archivo_excel}")
            raise SystemExit(1)
    else:
        archivo_excel = buscar_excel(carpeta_destino)

    print(f"\nExcel seleccionado:\n  {archivo_excel.name}")

    try:
        hoja, llaves = obtener_llaves(archivo_excel)
    except Exception as exc:
        print(f"\nERROR al leer el Excel: {exc}")
        raise SystemExit(1)

    print(f"Hoja encontrada: {hoja}")
    print(f"Llaves encontradas: {len(llaves)}\n")

    creadas, existentes, omitidas = crear_carpetas(llaves, carpeta_destino)

    print("\n" + "=" * 64)
    print("PROCESO TERMINADO")
    print(f"  Carpetas creadas:      {creadas}")
    print(f"  Ya existentes:         {existentes}")
    print(f"  Repetidas/omitidas:    {omitidas}")
    print("=" * 64)

    # Mantiene abierta la ventana si se ejecuta con doble clic.
    try:
        input("\nPresiona Enter para cerrar...")
    except EOFError:
        pass


if __name__ == "__main__":
    main()
