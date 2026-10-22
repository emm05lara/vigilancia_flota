"""Relación LLAVE <-> DNO (etapa 6B-1): solo lectura y validación.

                LLAVE
               /     \\
     carpeta local     DNO
          |             |
         GPX         BAANOOL

La LLAVE sigue siendo el identificador operativo; el DNO es únicamente el identificador técnico
para hablar con BAANOOL. Este módulo NO escribe en el Excel, NO descubre DNO por su cuenta y NO
deduce nada a partir de placas, modelo, LLAVE o nombres de archivo.

El DNO se trata como TEXTO (es un identificador, no un número). Lo que sea ambiguo se reporta para
revisión; solo se limpian los espacios externos. Un DNO vacío es SIN_DNO y no impide analizar los
GPX locales. Un DNO en estado REVISAR nunca queda disponible en `VehiculoBaanool.dno`.

No hay ninguna relación con el token de BAANOOL: este módulo no lo conoce.
"""

import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

from crear_carpetas_llaves import limpiar_nombre_carpeta, localizar_columna_llave

from .llaves import clave_carpeta

COLUMNA_DNO = "DNO"
PATRON_DNO = re.compile(r"^[A-Za-z0-9]+$")  # lo observado: dígitos; cualquier otra cosa se revisa
PATRON_NOTACION_CIENTIFICA = re.compile(r"^\d+([.,]\d+)?[eE][+-]?\d+$")
DIGITOS_EXACTOS_EXCEL = 15  # Excel guarda 15 cifras significativas en una celda numérica

# Estado del DNO de un vehículo
DNO_OK = "OK"
SIN_DNO = "SIN_DNO"
DNO_REVISAR = "REVISAR"

# Tipos de problema
PROBLEMA_SIN_COLUMNA_DNO = "SIN_COLUMNA_DNO"
PROBLEMA_DNO_INVALIDO = "DNO_INVALIDO"
PROBLEMA_DNO_DUPLICADO = "DNO_DUPLICADO"
PROBLEMA_LLAVE_DUPLICADA = "LLAVE_DUPLICADA"
PROBLEMA_CARPETA_DUPLICADA = "CARPETA_DUPLICADA"
PROBLEMA_DNO_CONFLICTO = "DNO_CONFLICTO"


@dataclass(frozen=True)
class VehiculoBaanool:
    llave_original: str
    nombre_carpeta: str  # nombre seguro para Windows, obtenido con limpiar_nombre_carpeta
    dno: str | None  # solo si es utilizable (estado OK); si no, None
    estado_dno: str  # DNO_OK | SIN_DNO | DNO_REVISAR
    dno_original: str | None = None  # el valor leído (texto), para diagnóstico
    observaciones: tuple[str, ...] = ()

    @property
    def listo_para_descarga(self) -> bool:
        return self.estado_dno == DNO_OK and self.dno is not None


@dataclass(frozen=True)
class ProblemaMapeo:
    tipo: str
    detalle: str
    llaves: tuple[str, ...] = ()
    fila: int | None = None  # fila del Excel (1 = primera)


@dataclass
class ResultadoMapeo:
    archivo: Path | None
    hoja: str | None
    tiene_columna_dno: bool
    vehiculos: list[VehiculoBaanool] = field(default_factory=list)
    problemas: list[ProblemaMapeo] = field(default_factory=list)

    @property
    def listos(self) -> list[VehiculoBaanool]:
        return [v for v in self.vehiculos if v.listo_para_descarga]

    @property
    def sin_dno(self) -> list[VehiculoBaanool]:
        return [v for v in self.vehiculos if v.estado_dno == SIN_DNO]

    @property
    def por_revisar(self) -> list[VehiculoBaanool]:
        return [v for v in self.vehiculos if v.estado_dno == DNO_REVISAR]

    def por_llave(self, llave: str) -> VehiculoBaanool | None:
        return next((v for v in self.vehiculos if v.llave_original == llave), None)

    def por_carpeta(self, nombre_carpeta: str) -> VehiculoBaanool | None:
        clave = clave_carpeta(nombre_carpeta)
        return next((v for v in self.vehiculos if clave_carpeta(v.nombre_carpeta) == clave), None)

    def resumen(self) -> dict[str, int]:
        """Solo conteos: no se imprimen listas de DNO."""
        por_tipo: dict[str, int] = {}
        for p in self.problemas:
            por_tipo[p.tipo] = por_tipo.get(p.tipo, 0) + 1
        return {"vehiculos": len(self.vehiculos), "listos_para_descarga": len(self.listos),
                "sin_dno": len(self.sin_dno), "por_revisar": len(self.por_revisar),
                "problemas": len(self.problemas), **{f"problema_{t}": n for t, n in por_tipo.items()}}


def enmascarar_dno(dno: str | None) -> str:
    """Para mensajes: solo los últimos 4 caracteres (no se listan DNO completos)."""
    if not dno:
        return "(vacío)"
    return "***" + dno[-4:] if len(dno) > 4 else "***"


@dataclass(frozen=True)
class DnoLeido:
    dno: str | None  # texto utilizable, o None
    estado: str
    original: str | None
    observaciones: tuple[str, ...] = ()


def _como_texto_original(valor) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def clasificar_dno(valor) -> DnoLeido:
    """Interpreta el valor de una celda de DNO. No inventa ni corrige más que espacios externos."""
    original = _como_texto_original(valor)

    if valor is None:
        return DnoLeido(None, SIN_DNO, None)
    if isinstance(valor, (bool, datetime, date, time)):
        return DnoLeido(None, DNO_REVISAR, original,
                        ("La celda no es texto ni número (tipo inesperado): no se puede usar como DNO",))

    if isinstance(valor, (int, float)):
        if isinstance(valor, float) and not valor.is_integer():
            return DnoLeido(None, DNO_REVISAR, original, ("Número con decimales: no es un identificador válido",))
        entero = int(valor)
        if entero < 0:
            return DnoLeido(None, DNO_REVISAR, original, ("Número negativo: no es un identificador válido",))
        digitos = str(entero)
        if len(digitos) > DIGITOS_EXACTOS_EXCEL:
            return DnoLeido(None, DNO_REVISAR, original,
                            (f"Celda numérica con {len(digitos)} dígitos: Excel solo conserva "
                             f"{DIGITOS_EXACTOS_EXCEL} cifras exactas, pudo perder dígitos. "
                             "Captura el DNO como texto",))
        return DnoLeido(digitos, DNO_OK, original,
                        ("Celda numérica: se leyó como texto; si el DNO tuviera ceros a la izquierda se habrían "
                         "perdido (se recomienda dar formato Texto a la columna)",))

    crudo = str(valor)
    limpio = crudo.strip().strip("﻿").strip()  # solo espacios externos
    if not limpio:
        return DnoLeido(None, SIN_DNO, None)
    observaciones = ("Se eliminaron espacios externos",) if limpio != crudo else ()
    if PATRON_NOTACION_CIENTIFICA.match(limpio):
        return DnoLeido(None, DNO_REVISAR, limpio,
                        observaciones + ("Parece notación científica: Excel pudo haber alterado el identificador",))
    if not PATRON_DNO.match(limpio):
        raros = sorted({c for c in limpio if not re.match(r"[A-Za-z0-9]", c)}, key=ord)
        return DnoLeido(None, DNO_REVISAR, limpio,
                        observaciones + ("Caracteres inesperados en el DNO: " +
                                         ", ".join(repr(c) for c in raros) + ". No se corrige automáticamente",))
    return DnoLeido(limpio, DNO_OK, limpio, observaciones)


def construir_mapeo(filas: list[tuple[int, str, object]], tiene_columna_dno: bool,
                    archivo: Path | None = None, hoja: str | None = None) -> ResultadoMapeo:
    """filas: (número de fila del Excel, LLAVE ya sin espacios externos, valor crudo de la celda DNO)."""
    resultado = ResultadoMapeo(archivo, hoja, tiene_columna_dno)
    if not tiene_columna_dno:
        resultado.problemas.append(ProblemaMapeo(
            PROBLEMA_SIN_COLUMNA_DNO,
            "El Excel no tiene columna DNO: el análisis local funciona, pero no está preparado para "
            "descarga por lote desde BAANOOL"))

    indice_carpeta: dict[str, int] = {}  # clave de carpeta -> posición en resultado.vehiculos
    for fila, llave, valor in filas:
        leido = clasificar_dno(valor) if tiene_columna_dno else DnoLeido(None, SIN_DNO, None)
        nombre = limpiar_nombre_carpeta(llave)
        clave = clave_carpeta(nombre)

        if clave in indice_carpeta:  # repetida: se conserva la primera aparición
            posicion = indice_carpeta[clave]
            existente = resultado.vehiculos[posicion]
            misma_llave = existente.llave_original.casefold() == llave.casefold()
            dnos_distintos = (leido.estado == DNO_OK and existente.estado_dno == DNO_OK
                              and leido.dno.casefold() != existente.dno.casefold())
            if misma_llave:
                resultado.problemas.append(ProblemaMapeo(
                    PROBLEMA_LLAVE_DUPLICADA, "La LLAVE aparece más de una vez; se usa la primera aparición",
                    (llave,), fila))
                if dnos_distintos:  # ninguno de los dos DNO es confiable
                    resultado.problemas.append(ProblemaMapeo(
                        PROBLEMA_DNO_CONFLICTO,
                        f"La misma LLAVE tiene DNO distintos ({enmascarar_dno(existente.dno)} y "
                        f"{enmascarar_dno(leido.dno)}); no se usa ninguno", (existente.llave_original,), fila))
                    resultado.vehiculos[posicion] = _marcar_revisar(
                        existente, "DNO distinto en una fila repetida de la misma LLAVE")
            else:  # LLAVEs distintas que producen la misma carpeta local
                detalle = (f"Dos LLAVEs distintas producen la misma carpeta {nombre!r}; solo la primera "
                           "tiene carpeta propia")
                if dnos_distintos:
                    detalle += (f" y sus DNO difieren ({enmascarar_dno(existente.dno)} y "
                                f"{enmascarar_dno(leido.dno)}): no se usa ninguno")
                    resultado.vehiculos[posicion] = _marcar_revisar(
                        existente, "Otra LLAVE comparte la misma carpeta con un DNO distinto")
                resultado.problemas.append(ProblemaMapeo(
                    PROBLEMA_CARPETA_DUPLICADA, detalle, (existente.llave_original, llave), fila))
            continue

        if leido.estado == DNO_REVISAR:
            resultado.problemas.append(ProblemaMapeo(
                PROBLEMA_DNO_INVALIDO, "; ".join(leido.observaciones), (llave,), fila))
        indice_carpeta[clave] = len(resultado.vehiculos)
        resultado.vehiculos.append(VehiculoBaanool(
            llave_original=llave, nombre_carpeta=nombre, dno=leido.dno, estado_dno=leido.estado,
            dno_original=leido.original, observaciones=leido.observaciones))

    _marcar_dno_duplicados(resultado)
    return resultado


def _marcar_revisar(v: VehiculoBaanool, motivo: str) -> VehiculoBaanool:
    return VehiculoBaanool(v.llave_original, v.nombre_carpeta, None, DNO_REVISAR, v.dno_original,
                           v.observaciones + (motivo,))


def _marcar_dno_duplicados(resultado: ResultadoMapeo) -> None:
    """Dos vehículos distintos con el mismo DNO: ninguno de los dos queda utilizable."""
    grupos: dict[str, list[int]] = {}
    for i, v in enumerate(resultado.vehiculos):
        if v.estado_dno == DNO_OK:
            grupos.setdefault(v.dno.casefold(), []).append(i)
    for posiciones in grupos.values():
        if len(posiciones) < 2:
            continue
        llaves = tuple(resultado.vehiculos[i].llave_original for i in posiciones)
        dno = resultado.vehiculos[posiciones[0]].dno
        resultado.problemas.append(ProblemaMapeo(
            PROBLEMA_DNO_DUPLICADO,
            f"El DNO {enmascarar_dno(dno)} está asignado a {len(posiciones)} LLAVEs distintas; no se usa "
            "para ninguna", llaves))
        for i in posiciones:
            resultado.vehiculos[i] = _marcar_revisar(resultado.vehiculos[i], "DNO duplicado en otra LLAVE")


def _celda(fila, posicion: int):
    try:
        return fila[posicion].value
    except IndexError:  # las filas de una hoja de solo lectura pueden ser más cortas
        return None


def leer_vehiculos(archivo_excel: Path) -> ResultadoMapeo:
    """Lee LLAVE y (si existe) DNO de la primera hoja con encabezado LLAVE. No modifica el archivo."""
    libro = load_workbook(archivo_excel, read_only=True, data_only=True)
    try:
        for hoja in libro.worksheets:
            posicion = localizar_columna_llave(hoja)
            if posicion is None:
                continue
            fila_encabezado, columna_llave = posicion

            encabezados = next(hoja.iter_rows(min_row=fila_encabezado, max_row=fila_encabezado), ())
            columna_dno = next((i for i, c in enumerate(encabezados, start=1)
                                if str(c.value).strip().upper() == COLUMNA_DNO), None)

            filas = []
            for numero, fila in enumerate(hoja.iter_rows(min_row=fila_encabezado + 1),
                                          start=fila_encabezado + 1):
                valor_llave = _celda(fila, columna_llave - 1)
                llave = "" if valor_llave is None else str(valor_llave).strip()
                if not llave:
                    continue
                filas.append((numero, llave, _celda(fila, columna_dno - 1) if columna_dno else None))
            return construir_mapeo(filas, columna_dno is not None, Path(archivo_excel), hoja.title)
    finally:
        libro.close()
    raise ValueError("No encontré una columna llamada 'LLAVE' en ninguna hoja del Excel.")
