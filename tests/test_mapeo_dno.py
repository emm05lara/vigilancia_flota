"""Pruebas del mapeo LLAVE <-> DNO (solo lectura y validación del Excel).

Los Excel se crean en un directorio temporal; Lista_Vehiculos.xlsx real no se toca.

    py -m unittest tests.test_mapeo_dno -v
"""

import hashlib
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from crear_carpetas_llaves import limpiar_nombre_carpeta, obtener_llaves
from src.llaves import cargar_llaves, indexar_por_carpeta, llave_de_carpeta
from src.mapeo_dno import (DNO_OK, DNO_REVISAR, PROBLEMA_CARPETA_DUPLICADA, PROBLEMA_DNO_CONFLICTO,
                           PROBLEMA_DNO_DUPLICADO, PROBLEMA_DNO_INVALIDO, PROBLEMA_LLAVE_DUPLICADA,
                           PROBLEMA_SIN_COLUMNA_DNO, SIN_DNO, clasificar_dno, enmascarar_dno,
                           leer_vehiculos)

L1, L2, L3 = "AVEO|2022|AZUL|PCH9635", "AVEO|2018|ARENA|F86AYT", "VENTO|2021|ROJO|SIN PLACAS"


def crear_excel(carpeta: Path, filas, con_dno=True, fila_encabezado=1, nombre="Lista.xlsx") -> Path:
    """filas: lista de (LLAVE, valor_DNO). El valor se escribe tal cual (str, int, None...)."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Hoja1"
    for i in range(1, fila_encabezado):
        hoja.cell(row=i, column=1, value="Lista de vehículos (título)")
    hoja.cell(row=fila_encabezado, column=1, value="LLAVE")
    if con_dno:
        hoja.cell(row=fila_encabezado, column=2, value="DNO")
    for i, (llave, dno) in enumerate(filas, start=fila_encabezado + 1):
        hoja.cell(row=i, column=1, value=llave)
        if con_dno and dno is not None:
            hoja.cell(row=i, column=2, value=dno)
    ruta = carpeta / nombre
    libro.save(ruta)
    return ruta


class PruebasMapeoDno(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.carpeta = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def leer(self, filas, **kwargs):
        return leer_vehiculos(crear_excel(self.carpeta, filas, **kwargs))

    # ---------- A, B, C ----------

    def test_A_llave_y_dno_validos(self):
        r = self.leer([(L1, "864035054122222"), (L2, "864035054133333")])
        self.assertTrue(r.tiene_columna_dno)
        v = r.por_llave(L1)
        self.assertEqual(v.llave_original, L1)
        self.assertEqual(v.nombre_carpeta, "AVEO - 2022 - AZUL - PCH9635")  # limpiar_nombre_carpeta
        self.assertEqual(v.nombre_carpeta, limpiar_nombre_carpeta(L1))
        self.assertEqual(v.dno, "864035054122222")
        self.assertEqual(v.estado_dno, DNO_OK)
        self.assertTrue(v.listo_para_descarga)
        self.assertEqual(r.problemas, [])
        self.assertEqual(r.hoja, "Hoja1")

    def test_B_dno_se_lee_como_texto(self):
        # Celda numérica (como la guarda Excel) y celda de texto con ceros a la izquierda.
        r = self.leer([(L1, 864035054122222), (L2, "000123456789012"), (L3, 864035054133333.0)])
        self.assertEqual(r.por_llave(L1).dno, "864035054122222")
        self.assertIsInstance(r.por_llave(L1).dno, str)
        self.assertEqual(r.por_llave(L2).dno, "000123456789012")  # el texto conserva los ceros
        self.assertEqual(r.por_llave(L3).dno, "864035054133333")  # sin ".0"
        self.assertTrue(any("numérica" in o for o in r.por_llave(L1).observaciones))  # se avisa, no es silencioso

    def test_C_dno_vacio_es_sin_dno_y_no_bloquea(self):
        r = self.leer([(L1, None), (L2, "   "), (L3, "864035054122222")])
        for llave in (L1, L2):
            v = r.por_llave(llave)
            self.assertEqual((v.estado_dno, v.dno), (SIN_DNO, None))
            self.assertFalse(v.listo_para_descarga)
        self.assertEqual(r.por_llave(L3).estado_dno, DNO_OK)
        self.assertEqual(len(r.vehiculos), 3)  # los vehículos sin DNO siguen existiendo (análisis local)
        self.assertEqual(r.problemas, [])  # vacío no es un error: es SIN_DNO
        self.assertEqual(len(r.sin_dno), 2)

    # ---------- D, E ----------

    def test_D_dos_llaves_con_el_mismo_dno(self):
        r = self.leer([(L1, "864035054122222"), (L2, "864035054122222"), (L3, "864035054133333")])
        for llave in (L1, L2):
            v = r.por_llave(llave)
            self.assertEqual((v.estado_dno, v.dno), (DNO_REVISAR, None))  # nunca queda utilizable
            self.assertEqual(v.dno_original, "864035054122222")  # se conserva para diagnóstico
        self.assertEqual(r.por_llave(L3).estado_dno, DNO_OK)
        dup = [p for p in r.problemas if p.tipo == PROBLEMA_DNO_DUPLICADO]
        self.assertEqual(len(dup), 1)
        self.assertEqual(set(dup[0].llaves), {L1, L2})
        self.assertNotIn("864035054122222", dup[0].detalle)  # el mensaje no lista el DNO completo
        self.assertIn("***2222", dup[0].detalle)

    def test_E_misma_llave_repetida(self):
        r = self.leer([(L1, "864035054122222"), (L2, "864035054133333"), (L1, "864035054122222")])
        self.assertEqual(len(r.vehiculos), 2)  # se conserva la primera aparición
        self.assertEqual(r.por_llave(L1).estado_dno, DNO_OK)  # mismo DNO: sigue utilizable
        tipos = [p.tipo for p in r.problemas]
        self.assertEqual(tipos, [PROBLEMA_LLAVE_DUPLICADA])
        self.assertEqual(r.problemas[0].fila, 4)

    def test_E2_llave_repetida_con_dno_distinto_es_conflicto(self):
        r = self.leer([(L1, "864035054122222"), (L1, "864035054199999")])
        self.assertEqual(len(r.vehiculos), 1)
        self.assertEqual((r.vehiculos[0].estado_dno, r.vehiculos[0].dno), (DNO_REVISAR, None))
        tipos = {p.tipo for p in r.problemas}
        self.assertEqual(tipos, {PROBLEMA_LLAVE_DUPLICADA, PROBLEMA_DNO_CONFLICTO})

    def test_E3_dos_llaves_distintas_con_la_misma_carpeta(self):
        r = self.leer([("A|B", "111111111111111"), ("A - B", "222222222222222")])
        self.assertEqual(len(r.vehiculos), 1)
        self.assertEqual([p.tipo for p in r.problemas], [PROBLEMA_CARPETA_DUPLICADA])
        self.assertEqual((r.vehiculos[0].estado_dno, r.vehiculos[0].dno), (DNO_REVISAR, None))  # DNO distintos
        # Con el mismo DNO (o sin DNO en la repetida) la primera sigue utilizable; se reporta igual.
        r = self.leer([("A|B", "111111111111111"), ("A - B", "111111111111111")])
        self.assertEqual(r.vehiculos[0].estado_dno, DNO_OK)
        self.assertEqual([p.tipo for p in r.problemas], [PROBLEMA_CARPETA_DUPLICADA])

    # ---------- F, G ----------

    def test_F_espacios_externos_se_limpian_y_se_avisa(self):
        r = self.leer([(L1, "  864035054122222 "), (L2, "\t864035054133333 ")])
        for llave, esperado in ((L1, "864035054122222"), (L2, "864035054133333")):
            v = r.por_llave(llave)
            self.assertEqual(v.dno, esperado)
            self.assertEqual(v.estado_dno, DNO_OK)
            self.assertIn("Se eliminaron espacios externos", v.observaciones)

    def test_F2_espacios_internos_no_se_corrigen(self):
        v = leer_vehiculos(crear_excel(self.carpeta, [(L1, "8640 35054122222")])).vehiculos[0]
        self.assertEqual((v.estado_dno, v.dno), (DNO_REVISAR, None))
        self.assertEqual(v.dno_original, "8640 35054122222")  # no se "arregla" quitando el espacio

    def test_G_caracteres_inesperados_se_reportan_sin_inventar(self):
        casos = ["86403-5054122", "ABC#123", "8.64035E+14", "864035054122222.5x", "86403505412222\n2"]
        for texto in casos:
            r = self.leer([(L1, texto)])
            v = r.vehiculos[0]
            self.assertEqual(v.estado_dno, DNO_REVISAR, texto)
            self.assertIsNone(v.dno, texto)
            self.assertEqual([p.tipo for p in r.problemas], [PROBLEMA_DNO_INVALIDO])

    def test_G2_numeros_ambiguos_de_excel(self):
        # 16 dígitos: Excel solo conserva 15 cifras exactas; decimales; fecha: ninguno se acepta.
        for valor in (8640350541222221, 864035054122222.5, datetime(2026, 10, 12), True, -5):
            v = self.leer([(L1, valor)]).vehiculos[0]
            self.assertEqual((v.estado_dno, v.dno), (DNO_REVISAR, None), repr(valor))
        self.assertEqual(clasificar_dno(999999999999999).dno, "999999999999999")  # 15 dígitos: exacto

    def test_notacion_cientifica_no_se_interpreta(self):
        d = clasificar_dno("8.64035E+14")
        self.assertEqual((d.estado, d.dno), (DNO_REVISAR, None))
        self.assertTrue(any("científica" in o for o in d.observaciones))

    # ---------- H, I ----------

    def test_H_excel_sin_columna_dno(self):
        ruta = crear_excel(self.carpeta, [(L1, None), (L2, None)], con_dno=False)
        r = leer_vehiculos(ruta)
        self.assertFalse(r.tiene_columna_dno)
        self.assertEqual(len(r.vehiculos), 2)
        self.assertTrue(all(v.estado_dno == SIN_DNO and not v.listo_para_descarga for v in r.vehiculos))
        self.assertEqual([p.tipo for p in r.problemas], [PROBLEMA_SIN_COLUMNA_DNO])
        self.assertEqual(r.listos, [])
        # El flujo local de LLAVE -> carpeta sigue funcionando con ese Excel.
        self.assertEqual(cargar_llaves(ruta), [L1, L2])
        indice = indexar_por_carpeta(cargar_llaves(ruta))
        self.assertEqual(llave_de_carpeta("aveo - 2022 - azul - pch9635", indice), L1)

    def test_I_las_funciones_de_llave_existentes_no_cambian_con_columna_dno(self):
        ruta = crear_excel(self.carpeta, [(L1, "864035054122222"), (L2, None), (L3, "x y")])
        hoja, llaves = obtener_llaves(ruta)  # la función original del script de carpetas
        self.assertEqual((hoja, llaves), ("Hoja1", [L1, L2, L3]))
        self.assertEqual(cargar_llaves(ruta), [L1, L2, L3])
        indice = indexar_por_carpeta(cargar_llaves(ruta))
        self.assertEqual(llave_de_carpeta(limpiar_nombre_carpeta(L3), indice), L3)

    def test_relacion_con_carpetas_y_encabezado_desplazado(self):
        ruta = crear_excel(self.carpeta, [(L1, "864035054122222")], fila_encabezado=3)
        r = leer_vehiculos(ruta)
        v = r.por_carpeta("AVEO - 2022 - AZUL - PCH9635")
        self.assertIs(v, r.por_llave(L1))
        self.assertEqual(v.dno, "864035054122222")
        self.assertIs(r.por_carpeta("aveo - 2022 - azul - pch9635"), v)  # Windows no distingue mayúsculas
        self.assertIsNone(r.por_carpeta("otra carpeta"))

    def test_leer_no_modifica_el_excel(self):
        ruta = crear_excel(self.carpeta, [(L1, "864035054122222"), (L2, None)])
        antes = (hashlib.sha256(ruta.read_bytes()).hexdigest(), ruta.stat().st_mtime_ns)
        leer_vehiculos(ruta)
        despues = (hashlib.sha256(ruta.read_bytes()).hexdigest(), ruta.stat().st_mtime_ns)
        self.assertEqual(antes, despues)

    def test_resumen_solo_con_conteos_y_mascara(self):
        r = self.leer([(L1, "864035054122222"), (L2, None), (L3, "bad id")])
        resumen = r.resumen()
        self.assertEqual((resumen["vehiculos"], resumen["listos_para_descarga"], resumen["sin_dno"],
                          resumen["por_revisar"]), (3, 1, 1, 1))
        self.assertTrue(all(isinstance(v, int) for v in resumen.values()))
        self.assertEqual(enmascarar_dno("864035054122222"), "***2222")
        self.assertEqual(enmascarar_dno(None), "(vacío)")


if __name__ == "__main__":
    unittest.main()
